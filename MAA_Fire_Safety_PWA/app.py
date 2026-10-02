#!/usr/bin/env python3
import base64, io, json, mimetypes, os, re, shutil, subprocess, tempfile, uuid
from datetime import date, datetime
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse, parse_qs, unquote

import requests
from docx import Document
from pypdf import PdfReader, PdfWriter

BASE = Path(__file__).resolve().parent
STATIC = BASE / 'static'
TEMPLATES = BASE / 'templates'
PORT = int(os.environ.get('PORT', '8080'))
SUPABASE_URL = os.environ.get('SUPABASE_URL', '').rstrip('/')
SUPABASE_SERVICE_KEY = os.environ.get('SUPABASE_SERVICE_ROLE_KEY', '')
ADMIN_PIN = os.environ.get('APP_ADMIN_PIN', '')
PHOTO_BUCKET = os.environ.get('SUPABASE_PHOTO_BUCKET', 'inspection-photos')
REMOTE = bool(SUPABASE_URL and SUPABASE_SERVICE_KEY)
APP_VERSION = '7.2-refined'

TESTS = {
    'extinguishers': {'name':'Fire Extinguishers','frequency':'Monthly','months':1,'template':'fire_extinguisher.docx','kind':'extinguisher','description':'Monthly extinguisher inspection','sort_order':1},
    'fire_doors': {'name':'Fire Doors','frequency':'Every 6 months','months':6,'template':'fire_door.docx','kind':'standard','tables':[(0,range(1,13))],'description':'Six-monthly fire door inspection','sort_order':2},
    'electronic_locks': {'name':'Electronic Locks','frequency':'Monthly','months':1,'template':'electronic_lock.docx','kind':'standard','tables':[(0,range(1,9))],'description':'Monthly electronic lock inspection','sort_order':3},
    'emergency_lighting': {'name':'Emergency Lighting','frequency':'Monthly','months':1,'template':'emergency_lighting.docx','kind':'standard','tables':[(0,range(1,15)),(1,range(0,10))],'description':'Monthly emergency lighting inspection','sort_order':4},
    'final_exits': {'name':'Final Exits','frequency':'Monthly','months':1,'template':'final_exit.docx','kind':'standard','tables':[(0,range(1,4))],'description':'Monthly final exit inspection','sort_order':5},
}

# ---------- local source parsing / seeding ----------
def source_locations(test_id):
    cfg = TESTS[test_id]
    doc = Document(TEMPLATES / cfg['template'])
    rows = []
    if cfg['kind'] == 'extinguisher':
        for ti, t in enumerate(doc.tables):
            for ri, row in enumerate(t.rows):
                loc = re.sub(r'\s+', ' ', row.cells[1].text.strip().replace('\n',' '))
                if ri == 0 or not loc or loc.lower() == 'location':
                    continue
                etype = 'H₂O' if loc.startswith('H₂O') else ('CO₂' if loc.startswith('CO₂') else '')
                rows.append({
                    'source_table': ti, 'source_row': ri, 'item_code': '', 'location': loc,
                    'extinguisher_type': etype, 'default_photo': f'/reference/extinguishers/ext_{len(rows)+1:02d}.jpg'
                })
    else:
        for ti, ris in cfg['tables']:
            t = doc.tables[ti]
            for ri in ris:
                row = t.rows[ri]
                rows.append({
                    'source_table': ti, 'source_row': ri,
                    'item_code': re.sub(r'\s+', ' ', row.cells[1].text.strip()),
                    'location': re.sub(r'\s+', ' ', row.cells[2].text.strip().replace('\n',' ')),
                    'extinguisher_type': '', 'default_photo': None
                })
    for i, x in enumerate(rows, 1):
        x['sort_order'] = i
    return rows

# ---------- Supabase REST helpers ----------
def sb_headers(extra=None):
    h = {'apikey': SUPABASE_SERVICE_KEY, 'Authorization': f'Bearer {SUPABASE_SERVICE_KEY}', 'Content-Type':'application/json'}
    if extra: h.update(extra)
    return h

def sb(table, method='GET', params=None, payload=None, prefer=None):
    if not REMOTE:
        raise RuntimeError('Supabase is not configured on Render yet.')
    headers = sb_headers({'Prefer': prefer}) if prefer else sb_headers()
    r = requests.request(method, f'{SUPABASE_URL}/rest/v1/{table}', headers=headers, params=params, json=payload, timeout=30)
    if r.status_code >= 300:
        raise RuntimeError(f'Supabase {table} error {r.status_code}: {r.text[:500]}')
    if not r.text:
        return None
    return r.json()

def ensure_seeded():
    if not REMOTE: return
    existing = sb('inspection_types', params={'select':'id','limit':'1'})
    if existing: return
    types = []
    for tid,c in TESTS.items():
        types.append({'id':tid,'name':c['name'],'frequency':c['frequency'],'months':c['months'],'description':c['description'],'sort_order':c['sort_order']})
    sb('inspection_types','POST',payload=types,prefer='return=minimal')
    items=[]
    for tid in TESTS:
        for x in source_locations(tid):
            items.append({'test_id':tid,'item_code':x['item_code'],'location':x['location'],'extinguisher_type':x['extinguisher_type'] or None,
                          'sort_order':x['sort_order'],'active':True,'source_table':x['source_table'],'source_row':x['source_row'],
                          'default_photo':x['default_photo'],'source_kind':'official'})
    sb('inspection_items','POST',payload=items,prefer='return=minimal')

def list_items(test_id, active_only=True):
    if not REMOTE:
        return [dict(id=f'local-{test_id}-{i}', test_id=test_id, item_code=x['item_code'], location=x['location'], extinguisher_type=x['extinguisher_type'],
                     sort_order=x['sort_order'], active=True, source_table=x['source_table'], source_row=x['source_row'], default_photo=x['default_photo'], photo_path=None, source_kind='official')
                for i,x in enumerate(source_locations(test_id),1)]
    ensure_seeded()
    params={'select':'*','test_id':f'eq.{test_id}','order':'sort_order.asc,created_at.asc'}
    if active_only: params['active']='eq.true'
    return sb('inspection_items',params=params)

def test_summaries():
    if not REMOTE:
        return [{'id':tid,'name':c['name'],'frequency':c['frequency'],'months':c['months'],'description':c['description'],'count':len(source_locations(tid)), 'mode':'local'} for tid,c in TESTS.items()]
    ensure_seeded()
    types = sb('inspection_types', params={'select':'*','order':'sort_order.asc'})
    out=[]
    for t in types:
        cnt = sb('inspection_items', params={'select':'id','test_id':f"eq.{t['id']}",'active':'eq.true'})
        last = sb('inspections', params={'select':'id,inspection_date,initials,status,completed_at','test_id':f"eq.{t['id']}",'status':'eq.completed','order':'inspection_date.desc','limit':'1'})
        prog = sb('inspections', params={'select':'id,inspection_date,initials,status,started_at','test_id':f"eq.{t['id']}",'status':'eq.in_progress','order':'started_at.desc','limit':'1'})
        prog_obj=prog[0] if prog else None
        if prog_obj:
            rr=sb('inspection_results',params={'select':'id,result','inspection_id':f"eq.{prog_obj['id']}"}) or []
            prog_obj={**prog_obj,'completed_checks':sum(1 for r in rr if r.get('result')),'total_checks':len(cnt or [])}
        out.append({**t,'count':len(cnt or []),'last_completed':last[0] if last else None,'in_progress':prog_obj,'mode':'supabase'})
    return out

# ---------- inspection lifecycle ----------
def start_or_resume(test_id, inspection_date, initials):
    ensure_seeded()
    existing = sb('inspections', params={'select':'*','test_id':f'eq.{test_id}','status':'eq.in_progress','order':'started_at.desc','limit':'1'})
    if existing:
        ins=existing[0]
    else:
        ins=sb('inspections','POST',payload={'test_id':test_id,'inspection_date':inspection_date,'initials':initials,'status':'in_progress'},prefer='return=representation')[0]
    items=list_items(test_id,True)
    rs=sb('inspection_results',params={'select':'*','inspection_id':f"eq.{ins['id']}"}) or []
    by={r['item_id']:r for r in rs}
    return {'inspection':ins,'items':items,'results':by}

def save_result(inspection_id,item_id,result,notes):
    if result not in ('','PASS','FAIL','ISSUE'): raise ValueError('Invalid result')
    payload={'inspection_id':inspection_id,'item_id':item_id,'result':result or None,'notes':notes or ''}
    params={'on_conflict':'inspection_id,item_id'}
    return sb('inspection_results','POST',params=params,payload=payload,prefer='resolution=merge-duplicates,return=representation')

# ---------- PDF generation ----------
def set_cell_text(cell,text):
    if cell.paragraphs:
        p=cell.paragraphs[0]
        if p.runs:
            p.runs[0].text=str(text or '')
            for r in p.runs[1:]: r.text=''
        else: p.add_run(str(text or ''))
        for extra in cell.paragraphs[1:]:
            for r in extra.runs: r.text=''
    else: cell.text=str(text or '')

def libreoffice_convert(docx_path, outdir):
    lo=shutil.which('libreoffice') or shutil.which('soffice')
    if not lo:
        mac='/Applications/LibreOffice.app/Contents/MacOS/soffice'
        if Path(mac).exists(): lo=mac
    if not lo: raise RuntimeError('LibreOffice is not installed on the server.')
    profile=Path(outdir)/('lo_profile_'+uuid.uuid4().hex); profile.mkdir()
    proc=subprocess.run([lo,'--headless',f'-env:UserInstallation=file://{profile}','--convert-to','pdf','--outdir',str(outdir),str(docx_path)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,timeout=120)
    pdf=Path(outdir)/(Path(docx_path).stem+'.pdf')
    if not pdf.exists(): raise RuntimeError('PDF conversion failed. '+proc.stderr[-500:])
    return pdf

def additional_doc(test_name, inspection_date, initials, extra_items, result_by):
    d=Document(); d.add_heading(f'{test_name} – Additional inspection items', level=1)
    p=d.add_paragraph(); p.add_run('Inspection date: ').bold=True; p.add_run(inspection_date)
    p.add_run('    Initials: ').bold=True; p.add_run(initials)
    table=d.add_table(rows=1, cols=5); table.style='Table Grid'
    hdr=table.rows[0].cells
    for c,txt in zip(hdr,['ID','Location','Type','Result','Notes']): c.text=txt
    for item in extra_items:
        r=result_by.get(item['id'],{})
        cells=table.add_row().cells
        vals=[item.get('item_code') or '', item.get('location') or '', item.get('extinguisher_type') or '', r.get('result') or '', r.get('notes') or '']
        for c,v in zip(cells,vals): c.text=str(v)
    return d

def make_pdf_remote(inspection_id):
    ins=sb('inspections',params={'select':'*','id':f'eq.{inspection_id}','limit':'1'})
    if not ins: raise ValueError('Inspection not found')
    ins=ins[0]; tid=ins['test_id']; cfg=TESTS[tid]
    items=list_items(tid,False)
    active=[x for x in items if x.get('active')]
    rs=sb('inspection_results',params={'select':'*','inspection_id':f'eq.{inspection_id}'}) or []
    by={r['item_id']:r for r in rs}
    missing=[x for x in active if not by.get(x['id'],{}).get('result')]
    if missing: raise ValueError(f'{len(missing)} active checks still need a result')
    inspection_date=ins['inspection_date']; initials=ins.get('initials') or 'F.G.'
    with tempfile.TemporaryDirectory() as td:
        td=Path(td); main_docx=td/f'{tid}_{inspection_date}.docx'; shutil.copy2(TEMPLATES/cfg['template'],main_docx)
        doc=Document(main_docx); y,m,d=inspection_date.split('-'); date_text=f'{y}/{m}/{d}'
        # clear/deactivate mapped rows and write all mapped active items using editable data
        for item in items:
            st,sr=item.get('source_table'),item.get('source_row')
            if st is None or sr is None: continue
            row=doc.tables[int(st)].rows[int(sr)]
            if cfg['kind']=='extinguisher':
                loc_col=1; id_col=None
            else:
                id_col=1; loc_col=2
            if not item.get('active'):
                set_cell_text(row.cells[0],'');
                if id_col is not None: set_cell_text(row.cells[id_col],'')
                set_cell_text(row.cells[loc_col],''); set_cell_text(row.cells[3],''); set_cell_text(row.cells[4],'')
                continue
            rr=by.get(item['id'],{})
            set_cell_text(row.cells[0],date_text)
            if id_col is not None: set_cell_text(row.cells[id_col],item.get('item_code') or '')
            set_cell_text(row.cells[loc_col],item.get('location') or '')
            set_cell_text(row.cells[3],rr.get('result') or '')
            set_cell_text(row.cells[4],initials)
        doc.save(main_docx); main_pdf=libreoffice_convert(main_docx,td)
        extras=[x for x in active if x.get('source_table') is None or x.get('source_row') is None]
        if not extras: return main_pdf.read_bytes()
        extra_doc=additional_doc(cfg['name'],inspection_date,initials,extras,by); extra_docx=td/'additional_items.docx'; extra_doc.save(extra_docx); extra_pdf=libreoffice_convert(extra_docx,td)
        writer=PdfWriter()
        for p in PdfReader(str(main_pdf)).pages: writer.add_page(p)
        for p in PdfReader(str(extra_pdf)).pages: writer.add_page(p)
        out=io.BytesIO(); writer.write(out); return out.getvalue()

def make_pdf_local(test_id,inspection_date,initials,results):
    # backward-compatible local mode
    cfg=TESTS[test_id]; rows=source_locations(test_id)
    if len(results)!=len(rows): raise ValueError(f'Expected {len(rows)} results, got {len(results)}')
    date.fromisoformat(inspection_date)
    with tempfile.TemporaryDirectory() as td:
        td=Path(td); out_docx=td/f'{test_id}_{inspection_date}.docx'; shutil.copy2(TEMPLATES/cfg['template'],out_docx)
        doc=Document(out_docx); y,m,d=inspection_date.split('-'); date_text=f'{y}/{m}/{d}'
        for meta,res in zip(rows,results):
            val=(res.get('result') or '').upper()
            if val not in ('PASS','FAIL','ISSUE'): raise ValueError('Every check must be PASS, FAIL, or ISSUE')
            row=doc.tables[meta['source_table']].rows[meta['source_row']]
            set_cell_text(row.cells[0],date_text); set_cell_text(row.cells[3],val); set_cell_text(row.cells[4],initials)
        doc.save(out_docx); return libreoffice_convert(out_docx,td).read_bytes()

# ---------- HTTP ----------
class Handler(SimpleHTTPRequestHandler):
    def log_message(self,fmt,*args): print('[MAA]',fmt%args)
    def json_body(self):
        n=int(self.headers.get('Content-Length','0')); return json.loads(self.rfile.read(n) or '{}')
    def send_json(self,obj,status=200):
        data=json.dumps(obj,ensure_ascii=False,default=str).encode(); self.send_response(status); self.send_header('Content-Type','application/json; charset=utf-8'); self.send_header('Content-Length',len(data)); self.send_header('Cache-Control','no-store'); self.end_headers(); self.wfile.write(data)
    def serve_file(self,path,ctype=None,cache='public, max-age=3600'):
        path=Path(path)
        if not path.exists() or not path.is_file(): return self.send_error(404)
        data=path.read_bytes(); self.send_response(200); self.send_header('Content-Type',ctype or mimetypes.guess_type(path.name)[0] or 'application/octet-stream'); self.send_header('Content-Length',len(data)); self.send_header('Cache-Control',cache); self.end_headers(); self.wfile.write(data)
    def admin_ok(self):
        return (not ADMIN_PIN) or self.headers.get('X-Admin-Pin','') == ADMIN_PIN
    def require_admin(self):
        if self.admin_ok(): return True
        self.send_json({'error':'Admin PIN required or incorrect'},403); return False
    def do_GET(self):
        try:
            p=urlparse(self.path).path
            if p=='/health': return self.send_json({'ok':True,'mode':'supabase' if REMOTE else 'local','admin_pin':bool(ADMIN_PIN),'version':APP_VERSION})
            if p=='/api/tests': return self.send_json({'today':date.today().isoformat(),'mode':'supabase' if REMOTE else 'local','tests':test_summaries()})
            if p=='/api/history':
                if not REMOTE: return self.send_json({'history':[],'mode':'local'})
                ensure_seeded(); hist=sb('inspections',params={'select':'id,test_id,inspection_date,initials,status,completed_at','status':'eq.completed','order':'inspection_date.desc,completed_at.desc','limit':'50'}) or []
                for h in hist: h['test_name']=TESTS.get(h['test_id'],{}).get('name',h['test_id'])
                # add outcome counts for the dashboard history
                for h in hist:
                    rr=sb('inspection_results',params={'select':'result','inspection_id':f"eq.{h['id']}"}) or []
                    h['counts']={k:sum(1 for x in rr if x.get('result')==k) for k in ('PASS','FAIL','ISSUE')}
                return self.send_json({'history':hist,'mode':'supabase'})
            m=re.fullmatch(r'/api/history/([^/]+)',p)
            if m:
                if not REMOTE: return self.send_json({'error':'History details require Supabase'},409)
                iid=m.group(1)
                ins=sb('inspections',params={'select':'id,test_id,inspection_date,initials,status,completed_at','id':f'eq.{iid}','limit':'1'}) or []
                if not ins: return self.send_json({'error':'Inspection not found'},404)
                rows=sb('inspection_results',params={'select':'item_id,result,notes,updated_at','inspection_id':f'eq.{iid}'}) or []
                items=sb('inspection_items',params={'select':'id,item_code,location,extinguisher_type,sort_order','test_id':f"eq.{ins[0]['test_id']}"}) or []
                by={x['id']:x for x in items}
                out=[]
                for r in rows:
                    it=by.get(r['item_id'],{})
                    out.append({**r,'item_code':it.get('item_code',''),'location':it.get('location','Unknown item'),'extinguisher_type':it.get('extinguisher_type'),'sort_order':it.get('sort_order',999)})
                out.sort(key=lambda x:(x.get('sort_order',999),x.get('location','')))
                z=ins[0]; z['test_name']=TESTS.get(z['test_id'],{}).get('name',z['test_id'])
                return self.send_json({'inspection':z,'results':out})
            m=re.fullmatch(r'/api/evidence/([^/]+)/([^/]+)/photo',p)
            if m:
                if not REMOTE: return self.send_error(404)
                inspection_id,item_id=m.group(1),m.group(2)
                path=f'evidence/{inspection_id}/{item_id}.jpg'
                r=requests.get(f'{SUPABASE_URL}/storage/v1/object/{PHOTO_BUCKET}/{path}',headers={'apikey':SUPABASE_SERVICE_KEY,'Authorization':f'Bearer {SUPABASE_SERVICE_KEY}'},timeout=30)
                if r.status_code>=300: return self.send_error(404)
                self.send_response(200); self.send_header('Content-Type',r.headers.get('Content-Type','image/jpeg')); self.send_header('Content-Length',len(r.content)); self.send_header('Cache-Control','no-store'); self.end_headers(); self.wfile.write(r.content); return
            if p=='/api/backup':
                if not self.require_admin(): return
                if not REMOTE: return self.send_json({'error':'Backup requires Supabase'},409)
                payload={'version':APP_VERSION,'exported_at':datetime.utcnow().isoformat()+'Z'}
                for table in ('inspection_types','inspection_items','inspections','inspection_results'):
                    payload[table]=sb(table,params={'select':'*'}) or []
                data=json.dumps(payload,ensure_ascii=False,indent=2,default=str).encode()
                self.send_response(200); self.send_header('Content-Type','application/json; charset=utf-8'); self.send_header('Content-Disposition','attachment; filename=MAA_Fire_Safety_Backup.json'); self.send_header('Content-Length',len(data)); self.send_header('Cache-Control','no-store'); self.end_headers(); self.wfile.write(data); return
            if p.startswith('/api/test/'):
                tid=p.split('/')[-1]
                if tid not in TESTS: return self.send_json({'error':'Unknown test'},404)
                c=TESTS[tid]; return self.send_json({'id':tid,'name':c['name'],'frequency':c['frequency'],'months':c['months'],'description':c['description'],'locations':list_items(tid,True),'mode':'supabase' if REMOTE else 'local'})
            if p.startswith('/api/manage/items/'):
                if not self.require_admin(): return
                tid=p.split('/')[-1]
                if tid not in TESTS: return self.send_json({'error':'Unknown test'},404)
                return self.send_json({'items':list_items(tid,False)})
            if p.startswith('/api/item/') and p.endswith('/photo'):
                item_id=p.split('/')[3]
                if not REMOTE: return self.send_error(404)
                rows=sb('inspection_items',params={'select':'photo_path','id':f'eq.{item_id}','limit':'1'})
                if not rows or not rows[0].get('photo_path'): return self.send_error(404)
                path=rows[0]['photo_path']
                r=requests.get(f'{SUPABASE_URL}/storage/v1/object/{PHOTO_BUCKET}/{path}',headers={'apikey':SUPABASE_SERVICE_KEY,'Authorization':f'Bearer {SUPABASE_SERVICE_KEY}'},timeout=30)
                if r.status_code>=300: return self.send_error(404)
                self.send_response(200); self.send_header('Content-Type',r.headers.get('Content-Type','image/jpeg')); self.send_header('Content-Length',len(r.content)); self.send_header('Cache-Control','no-store'); self.end_headers(); self.wfile.write(r.content); return
            if p=='/manifest.webmanifest': return self.serve_file(STATIC/'manifest.webmanifest','application/manifest+json','no-cache')
            if p=='/sw.js': return self.serve_file(STATIC/'sw.js','application/javascript','no-cache')
            if p.startswith('/icons/') or p.startswith('/reference/'): return self.serve_file(STATIC/p.lstrip('/'))
            return self.serve_file(STATIC/'index.html','text/html; charset=utf-8','no-cache')
        except Exception as e: self.send_json({'error':str(e)},500)

    def do_POST(self):
        try:
            p=urlparse(self.path).path
            if p=='/api/inspections/start':
                if not REMOTE: return self.send_json({'error':'Shared sync requires Supabase configuration'},409)
                b=self.json_body(); tid=b.get('test_id')
                if tid not in TESTS: raise ValueError('Unknown test')
                return self.send_json(start_or_resume(tid,b.get('inspection_date') or date.today().isoformat(),(b.get('initials') or 'F.G.').strip()))
            m=re.fullmatch(r'/api/inspections/([^/]+)/result',p)
            if m:
                b=self.json_body(); return self.send_json(save_result(m.group(1),b['item_id'],b.get('result',''),b.get('notes','')))
            m=re.fullmatch(r'/api/inspections/([^/]+)/complete',p)
            if m:
                if not REMOTE: return self.send_json({'error':'Shared sync requires Supabase configuration'},409)
                iid=m.group(1); b=self.json_body();
                # update initials/date before generation
                sb('inspections','PATCH',params={'id':f'eq.{iid}'},payload={'inspection_date':b.get('inspection_date'),'initials':b.get('initials')},prefer='return=minimal')
                data=make_pdf_remote(iid)
                sb('inspections','PATCH',params={'id':f'eq.{iid}'},payload={'status':'completed','completed_at':datetime.utcnow().isoformat()+'Z'},prefer='return=minimal')
                ins=sb('inspections',params={'select':'test_id,inspection_date','id':f'eq.{iid}','limit':'1'})[0]
                friendly=TESTS[ins['test_id']]['name'].replace(' ','_'); name=f"MAA_{friendly}_{ins['inspection_date']}.pdf"
                self.send_response(200); self.send_header('Content-Type','application/pdf'); self.send_header('Content-Disposition',f'attachment; filename="{name}"'); self.send_header('Content-Length',len(data)); self.send_header('Cache-Control','no-store'); self.end_headers(); self.wfile.write(data); return
            if p=='/api/generate':
                # local-mode legacy endpoint
                b=self.json_body(); tid=b['test_id']; data=make_pdf_local(tid,b.get('inspection_date') or date.today().isoformat(),(b.get('initials') or 'F.G.').strip(),b['results'])
                friendly=TESTS[tid]['name'].replace(' ','_'); name=f"MAA_{friendly}_{b.get('inspection_date')}.pdf"; self.send_response(200); self.send_header('Content-Type','application/pdf'); self.send_header('Content-Disposition',f'attachment; filename="{name}"'); self.send_header('Content-Length',len(data)); self.end_headers(); self.wfile.write(data); return
            m=re.fullmatch(r'/api/evidence/([^/]+)/([^/]+)/photo',p)
            if m:
                if not REMOTE: return self.send_json({'error':'Evidence photo storage requires Supabase configuration'},409)
                inspection_id,item_id=m.group(1),m.group(2)
                # Verify this item belongs to this inspection before accepting evidence.
                ins=sb('inspections',params={'select':'test_id','id':f'eq.{inspection_id}','limit':'1'}) or []
                item=sb('inspection_items',params={'select':'test_id','id':f'eq.{item_id}','limit':'1'}) or []
                if not ins or not item or ins[0]['test_id']!=item[0]['test_id']: return self.send_json({'error':'Inspection item mismatch'},409)
                b=self.json_body(); data=b.get('data_url',''); mm=re.match(r'data:(image/[^;]+);base64,(.+)',data)
                if not mm: raise ValueError('Invalid image')
                raw=base64.b64decode(mm.group(2))
                if len(raw)>3_000_000: raise ValueError('Photo is too large after compression')
                path=f'evidence/{inspection_id}/{item_id}.jpg'
                r=requests.post(f'{SUPABASE_URL}/storage/v1/object/{PHOTO_BUCKET}/{path}',headers={'apikey':SUPABASE_SERVICE_KEY,'Authorization':f'Bearer {SUPABASE_SERVICE_KEY}','Content-Type':'image/jpeg','x-upsert':'true'},data=raw,timeout=30)
                if r.status_code>=300: raise RuntimeError('Evidence photo upload failed: '+r.text[:300])
                return self.send_json({'ok':True,'url':f'/api/evidence/{inspection_id}/{item_id}/photo?t={int(datetime.utcnow().timestamp())}'})
            if p=='/api/manage/item':
                if not self.require_admin(): return
                if not REMOTE: return self.send_json({'error':'Item editing requires Supabase configuration'},409)
                b=self.json_body(); tid=b.get('test_id');
                if tid not in TESTS: raise ValueError('Unknown test')
                loc=(b.get('location') or '').strip()
                if not loc: raise ValueError('Location is required')
                payload={'test_id':tid,'item_code':(b.get('item_code') or '').strip(),'location':loc,'extinguisher_type':(b.get('extinguisher_type') or '').strip() or None,'sort_order':int(b.get('sort_order') or 999),'active':bool(b.get('active',True))}
                if b.get('id'):
                    # Preserve official/source metadata when editing an existing item.
                    out=sb('inspection_items','PATCH',params={'id':f"eq.{b['id']}"},payload=payload,prefer='return=representation')
                else:
                    payload['source_kind']='added'
                    out=sb('inspection_items','POST',payload=payload,prefer='return=representation')
                return self.send_json(out[0])
            m=re.fullmatch(r'/api/manage/item/([^/]+)/deactivate',p)
            if m:
                if not self.require_admin(): return
                out=sb('inspection_items','PATCH',params={'id':f'eq.{m.group(1)}'},payload={'active':False},prefer='return=representation'); return self.send_json(out[0])
            m=re.fullmatch(r'/api/manage/item/([^/]+)/activate',p)
            if m:
                if not self.require_admin(): return
                out=sb('inspection_items','PATCH',params={'id':f'eq.{m.group(1)}'},payload={'active':True},prefer='return=representation'); return self.send_json(out[0])
            m=re.fullmatch(r'/api/manage/item/([^/]+)/delete',p)
            if m:
                if not self.require_admin(): return
                iid=m.group(1); item=sb('inspection_items',params={'select':'source_kind','id':f'eq.{iid}','limit':'1'})
                if not item: return self.send_json({'error':'Item not found'},404)
                if item[0].get('source_kind')=='official': return self.send_json({'error':'Official items cannot be deleted. Deactivate them instead.'},409)
                used=sb('inspection_results',params={'select':'id','item_id':f'eq.{iid}','limit':'1'})
                if used: return self.send_json({'error':'This item has inspection history. Deactivate it instead.'},409)
                sb('inspection_items','DELETE',params={'id':f'eq.{iid}'},prefer='return=minimal'); return self.send_json({'ok':True})
            m=re.fullmatch(r'/api/item/([^/]+)/photo',p)
            if m:
                if not REMOTE: return self.send_json({'error':'Shared photo storage requires Supabase configuration'},409)
                b=self.json_body(); data=b.get('data_url',''); mm=re.match(r'data:(image/[^;]+);base64,(.+)',data)
                if not mm: raise ValueError('Invalid image')
                raw=base64.b64decode(mm.group(2));
                if len(raw)>3_000_000: raise ValueError('Photo is too large after compression')
                iid=m.group(1); ext='jpg' if 'jpeg' in mm.group(1) else ('png' if 'png' in mm.group(1) else 'webp'); path=f'items/{iid}.{ext}'
                r=requests.post(f'{SUPABASE_URL}/storage/v1/object/{PHOTO_BUCKET}/{path}',headers={'apikey':SUPABASE_SERVICE_KEY,'Authorization':f'Bearer {SUPABASE_SERVICE_KEY}','Content-Type':mm.group(1),'x-upsert':'true'},data=raw,timeout=30)
                if r.status_code>=300: raise RuntimeError('Photo upload failed: '+r.text[:300])
                sb('inspection_items','PATCH',params={'id':f'eq.{iid}'},payload={'photo_path':path},prefer='return=minimal'); return self.send_json({'ok':True,'url':f'/api/item/{iid}/photo?t={int(datetime.utcnow().timestamp())}'})
            return self.send_json({'error':'Not found'},404)
        except Exception as e: self.send_json({'error':str(e)},400)

    def do_DELETE(self):
        try:
            p=urlparse(self.path).path
            m=re.fullmatch(r'/api/evidence/([^/]+)/([^/]+)/photo',p)
            if m:
                if not REMOTE: return self.send_json({'error':'Evidence photo storage requires Supabase configuration'},409)
                path=f'evidence/{m.group(1)}/{m.group(2)}.jpg'
                requests.delete(f'{SUPABASE_URL}/storage/v1/object/{PHOTO_BUCKET}/{path}',headers={'apikey':SUPABASE_SERVICE_KEY,'Authorization':f'Bearer {SUPABASE_SERVICE_KEY}'},timeout=30)
                return self.send_json({'ok':True})
            m=re.fullmatch(r'/api/item/([^/]+)/photo',p)
            if m:
                if not REMOTE: return self.send_json({'error':'Shared photo storage requires Supabase configuration'},409)
                iid=m.group(1); rows=sb('inspection_items',params={'select':'photo_path','id':f'eq.{iid}','limit':'1'})
                if rows and rows[0].get('photo_path'):
                    path=rows[0]['photo_path']
                    requests.delete(f'{SUPABASE_URL}/storage/v1/object/{PHOTO_BUCKET}/{path}',headers={'apikey':SUPABASE_SERVICE_KEY,'Authorization':f'Bearer {SUPABASE_SERVICE_KEY}'},timeout=30)
                    sb('inspection_items','PATCH',params={'id':f'eq.{iid}'},payload={'photo_path':None},prefer='return=minimal')
                return self.send_json({'ok':True})
            return self.send_json({'error':'Not found'},404)
        except Exception as e: self.send_json({'error':str(e)},400)

if __name__=='__main__':
    print(f'MAA Fire Safety Checks listening on 0.0.0.0:{PORT} mode={"supabase" if REMOTE else "local"}')
    ThreadingHTTPServer(('0.0.0.0',PORT),Handler).serve_forever()
