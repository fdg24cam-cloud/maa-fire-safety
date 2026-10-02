#!/usr/bin/env python3
import io, json, mimetypes, os, re, shutil, subprocess, tempfile
from datetime import date
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse
from docx import Document

BASE = Path(__file__).resolve().parent
STATIC = BASE / 'static'
TEMPLATES = BASE / 'templates'
PORT = int(os.environ.get('PORT','8080'))

TESTS = {
    'extinguishers': {'name':'Fire Extinguishers','frequency':'Monthly','months':1,'template':'fire_extinguisher.docx','kind':'extinguisher','description':'Monthly extinguisher inspection'},
    'fire_doors': {'name':'Fire Doors','frequency':'Every 6 months','months':6,'template':'fire_door.docx','kind':'standard','tables':[(0,range(1,13))],'description':'Six-monthly fire door inspection'},
    'electronic_locks': {'name':'Electronic Locks','frequency':'Monthly','months':1,'template':'electronic_lock.docx','kind':'standard','tables':[(0,range(1,9))],'description':'Monthly electronic lock inspection'},
    'emergency_lighting': {'name':'Emergency Lighting','frequency':'Monthly','months':1,'template':'emergency_lighting.docx','kind':'standard','tables':[(0,range(1,15)),(1,range(0,10))],'description':'Monthly emergency lighting inspection'},
    'final_exits': {'name':'Final Exits','frequency':'Monthly','months':1,'template':'final_exit.docx','kind':'standard','tables':[(0,range(1,4))],'description':'Monthly final exit inspection'},
}

def locations_for(test_id):
    cfg=TESTS[test_id]; doc=Document(TEMPLATES/cfg['template']); rows=[]
    if cfg['kind']=='extinguisher':
        for ti,t in enumerate(doc.tables):
            for ri,row in enumerate(t.rows):
                loc=row.cells[1].text.strip().replace('\n',' ')
                if ri==0 or loc.lower()=='location': continue
                photo=f'/reference/extinguishers/ext_{len(rows)+1:02d}.jpg'
                rows.append({'table':ti,'row':ri,'id':'','location':re.sub(r'\s+',' ',loc),'reference_photo':photo})
    else:
        for ti,ris in cfg['tables']:
            t=doc.tables[ti]
            for ri in ris:
                row=t.rows[ri]
                rows.append({'table':ti,'row':ri,'id':re.sub(r'\s+',' ',row.cells[1].text.strip()),'location':re.sub(r'\s+',' ',row.cells[2].text.strip().replace('\n',' ')),'reference_photo':None})
    return rows

def set_cell_text(cell,text):
    if cell.paragraphs:
        p=cell.paragraphs[0]
        if p.runs:
            p.runs[0].text=text
            for r in p.runs[1:]: r.text=''
        else: p.add_run(text)
        for extra in cell.paragraphs[1:]:
            for r in extra.runs: r.text=''
    else: cell.text=text

def make_pdf(test_id,inspection_date,initials,results):
    cfg=TESTS[test_id]; rows=locations_for(test_id)
    if len(results)!=len(rows): raise ValueError(f'Expected {len(rows)} results, got {len(results)}')
    date.fromisoformat(inspection_date)
    with tempfile.TemporaryDirectory() as td:
        td=Path(td); out_docx=td/f'{test_id}_{inspection_date}.docx'; shutil.copy2(TEMPLATES/cfg['template'],out_docx)
        doc=Document(out_docx); y,m,d=inspection_date.split('-'); date_text=f'{y}/{m}/{d}'
        for meta,res in zip(rows,results):
            val=(res.get('result') or '').upper()
            if val not in ('PASS','FAIL','ISSUE'): raise ValueError('Every check must be PASS, FAIL, or ISSUE')
            row=doc.tables[meta['table']].rows[meta['row']]
            set_cell_text(row.cells[0],date_text); set_cell_text(row.cells[3],val); set_cell_text(row.cells[4],initials)
        doc.save(out_docx)
        lo=shutil.which('libreoffice') or shutil.which('soffice')
        if not lo:
            mac='/Applications/LibreOffice.app/Contents/MacOS/soffice'
            if Path(mac).exists(): lo=mac
        if not lo: raise RuntimeError('LibreOffice is not installed on the server.')
        profile=td/'lo_profile'; profile.mkdir()
        proc=subprocess.run([lo,'--headless',f'-env:UserInstallation=file://{profile}','--convert-to','pdf','--outdir',str(td),str(out_docx)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True,timeout=90)
        pdf=td/(out_docx.stem+'.pdf')
        if not pdf.exists(): raise RuntimeError('PDF conversion failed. '+proc.stderr[-500:])
        return pdf.read_bytes()

class Handler(SimpleHTTPRequestHandler):
    def log_message(self,fmt,*args): print('[MAA]',fmt%args)
    def send_json(self,obj,status=200):
        data=json.dumps(obj,ensure_ascii=False).encode(); self.send_response(status); self.send_header('Content-Type','application/json; charset=utf-8'); self.send_header('Content-Length',len(data)); self.send_header('Cache-Control','no-store'); self.end_headers(); self.wfile.write(data)
    def serve_file(self,path,ctype=None):
        path=Path(path)
        if not path.exists() or not path.is_file(): return self.send_error(404)
        data=path.read_bytes(); self.send_response(200); self.send_header('Content-Type',ctype or mimetypes.guess_type(path.name)[0] or 'application/octet-stream'); self.send_header('Content-Length',len(data)); self.end_headers(); self.wfile.write(data)
    def do_GET(self):
        p=urlparse(self.path).path
        if p=='/health': return self.send_json({'ok':True})
        if p=='/api/tests': return self.send_json({'today':date.today().isoformat(),'tests':[{'id':tid,'name':c['name'],'frequency':c['frequency'],'months':c['months'],'description':c['description'],'count':len(locations_for(tid))} for tid,c in TESTS.items()]})
        if p.startswith('/api/test/'):
            tid=p.split('/')[-1]
            if tid not in TESTS: return self.send_json({'error':'Unknown test'},404)
            c=TESTS[tid]; return self.send_json({'id':tid,'name':c['name'],'frequency':c['frequency'],'months':c['months'],'description':c['description'],'locations':locations_for(tid)})
        if p=='/manifest.webmanifest': return self.serve_file(STATIC/'manifest.webmanifest','application/manifest+json')
        if p=='/sw.js': return self.serve_file(STATIC/'sw.js','application/javascript')
        if p.startswith('/icons/') or p.startswith('/reference/'): return self.serve_file(STATIC/p.lstrip('/'))
        return self.serve_file(STATIC/'index.html','text/html; charset=utf-8')
    def do_POST(self):
        p=urlparse(self.path).path
        if p!='/api/generate': return self.send_json({'error':'Not found'},404)
        try:
            n=int(self.headers.get('Content-Length','0')); body=json.loads(self.rfile.read(n) or '{}')
            tid=body['test_id']
            if tid not in TESTS: raise ValueError('Unknown test')
            inspection_date=body.get('inspection_date') or date.today().isoformat(); initials=(body.get('initials') or 'F.G.').strip()
            if not initials: raise ValueError('Initials are required')
            data=make_pdf(tid,inspection_date,initials,body['results'])
            friendly=TESTS[tid]['name'].replace(' ','_'); name=f'MAA_{friendly}_{inspection_date}.pdf'
            self.send_response(200); self.send_header('Content-Type','application/pdf'); self.send_header('Content-Disposition',f'attachment; filename="{name}"'); self.send_header('Content-Length',len(data)); self.send_header('Cache-Control','no-store'); self.end_headers(); self.wfile.write(data)
        except Exception as e: self.send_json({'error':str(e)},400)

if __name__=='__main__':
    print(f'MAA Fire Safety Checks listening on 0.0.0.0:{PORT}')
    ThreadingHTTPServer(('0.0.0.0',PORT),Handler).serve_forever()
